"""가져온 코드의 규격을 실제 문서로 확인한다. 모델 가중치는 받지 않는다."""
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock

import numpy as np
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from pptx import Presentation
from pptx.util import Inches

from file_loader.document_loader import DocumentLoader, DocumentScanner, DocumentChunk
from similarity.document_embedder import DocumentEmbedder
from file_encoder.document_encoder import DocumentEncoder


class ImportedDocumentTests(unittest.TestCase):
    def test_pdf_page_text_and_number(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'test.pdf'
            writer = PdfWriter()
            page = writer.add_blank_page(width=300, height=300)
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                     NameObject('/Subtype'): NameObject('/Type1'),
                                     NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
            content = DecodedStreamObject()
            content.set_data(b'BT /F1 12 Tf 20 200 Td (Document integration check) Tj ET')
            page[NameObject('/Contents')] = writer._add_object(content)
            with path.open('wb') as output:
                writer.write(output)
            chunks = DocumentLoader().process_file(str(path))
            self.assertEqual(len(chunks), 1)
            self.assertEqual(chunks[0].page_number, 1)
            self.assertIn('Document integration check', chunks[0].text)
            self.assertEqual(chunks[0].chunk_type, 'pdf')

    def test_pptx_slide_text_and_number(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'slides.pptx'
            presentation = Presentation()
            for text in ('First document page', 'Second document page'):
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = text
            presentation.save(path)
            chunks = DocumentLoader().process_file(str(path))
            self.assertEqual([c.page_number for c in chunks], [1, 2])
            self.assertEqual([c.text for c in chunks], ['First document page', 'Second document page'])

    def test_image_chunk_and_scanner(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'photo.png'
            Image.new('RGB', (8, 8), 'white').save(path)
            (Path(folder) / 'unsupported.txt').write_text('text', encoding='utf-8')
            self.assertEqual(list(DocumentScanner(folder).scan_files()), [str(path)])
            chunk = DocumentLoader().process_file(str(path))[0]
            self.assertEqual(chunk.chunk_type, 'image')
            self.assertEqual(chunk.text, '')
            self.assertEqual(chunk.image.mode, 'RGB')

    def test_broken_document_returns_empty_list(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'broken.pdf'
            path.write_bytes(b'broken document')
            self.assertEqual(DocumentLoader().process_file(str(path)), [])

    def test_embedder_search_uses_its_own_model(self):
        # 인코더 출력만 테스트 대역이다. 검색 수식은 가져온 원본을 그대로 실행한다.
        embedder = object.__new__(DocumentEmbedder)
        embedder.model = Mock()
        embedder.model.encode.side_effect = [np.array([[1., 0.], [0., 1.]]), np.array([[0., 1.]])]
        chunks = [DocumentChunk('a.pdf', 'a.pdf', 1, 'first', 'pdf'),
                  DocumentChunk('b.pdf', 'b.pdf', 2, 'second', 'pdf')]
        embedder.embed_chunks(chunks)
        results = embedder.search('second query', top_k=1)
        self.assertIs(results[0]['chunk'], chunks[1])
        embedder.model.encode.assert_called_with(['second query'])

    def test_image_encoder_wrapper_forwards_pil_input(self):
        image_encoder = Mock()
        encoder = DocumentEncoder(image_encoder=image_encoder)
        image = Image.new('RGB', (8, 8))
        encoder.encode(image)
        image_encoder.encode.assert_called_once_with(image)
        encoder.encode_batch([image])
        image_encoder.encode_batch.assert_called_once_with([image])
