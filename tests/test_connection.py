"""실제 로더·SQLite와 팀원 검색 수식으로 연결을 검증한다. 모델 출력만 대역이다."""
from io import BytesIO
from pathlib import Path
import gc
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock

import numpy as np
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from pptx import Presentation
from pptx.util import Inches

from connection.service import DocumentSearch
from database.db_manager import Repository
from file_loader.document_loader import DocumentLoader
from file_encoder.document_encoder import DocumentEncoder
from similarity.document_embedder import DocumentEmbedder


class ModelOutput:
    def encode(self, texts):
        return np.array([[1., 0.] if 'apple' in text else [0., 1.] for text in texts])


class ConnectionTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.addCleanup(gc.collect)  # 팀원 DB의 원본 연결은 GC 전까지 남을 수 있다.
        self.folder = Path(self.directory.name)
        self.image_encoder = Mock()
        self.image_encoder.encode.return_value = np.array([0., 1.])
        self.document_encoder = Mock(wraps=DocumentEncoder(self.image_encoder))
        self.query_encoder = Mock()
        self.query_encoder.encode_text.return_value = np.array([0., 1.])
        self.embedder = object.__new__(DocumentEmbedder)
        self.embedder.model = Mock(wraps=ModelOutput())
        self.embedder.chunks, self.embedder.embeddings = [], None
        self.repository = Repository(str(self.folder / 'metadata.db'))
        self.service = DocumentSearch(DocumentLoader(), self.image_encoder, self.document_encoder,
                                      self.embedder, self.query_encoder, self.repository)

    def make_documents(self):
        writer = PdfWriter()
        for text in ('apple first page', 'apple second page'):
            page = writer.add_blank_page(width=300, height=300)
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
            stream = DecodedStreamObject()
            stream.set_data(f'BT /F1 12 Tf 20 200 Td ({text}) Tj ET'.encode())
            page[NameObject('/Contents')] = writer._add_object(stream)
        with (self.folder / 'apple.pdf').open('wb') as output:
            writer.write(output)
        picture = Image.new('RGB', (12, 12), 'yellow')
        picture.save(self.folder / 'photo.png')
        memory = BytesIO()
        picture.save(memory, format='PNG')
        memory.seek(0)
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = 'banana slide'
        slide.shapes.add_picture(memory, Inches(1), Inches(3), width=Inches(1))
        presentation.save(self.folder / 'banana.pptx')

    def test_routes_loaders_and_encoders_and_saves_each_file_once(self):
        self.make_documents()
        documents, failures = self.service.analyze(self.folder)
        self.assertEqual({d.file_name for d in documents}, {'apple.pdf', 'banana.pptx', 'photo.png'})
        self.assertEqual(failures, [])
        self.assertEqual(len(self.service.file_ids), 3)
        self.assertEqual(len(self.service.text_chunks), 3)
        self.document_encoder.encode.assert_called_once()
        self.assertEqual(self.image_encoder.encode.call_count, 2)
        self.assertEqual(len(self.service.image_vectors), 2)
        for file_id in self.service.file_ids.values():
            self.assertIsNotNone(self.repository.find_by_id(file_id))

    def test_document_search_deduplicates_pages_and_does_not_use_clip_query(self):
        self.make_documents()
        self.service.analyze(self.folder)
        results = self.service.search('apple query', mode='documents')
        self.assertEqual(results[0][0].file_name, 'apple.pdf')
        self.assertEqual(len([r for r in results if r[0].file_name == 'apple.pdf']), 1)
        self.query_encoder.encode_text.assert_not_called()
        self.embedder.model.encode.assert_called_with(['apple query'])

    def test_image_search_uses_clip_query_and_returns_file_info(self):
        self.make_documents()
        self.service.analyze(self.folder)
        self.embedder.model.encode.reset_mock()
        results = self.service.search('yellow image', mode='images')
        self.assertEqual({info.file_name for info, score in results}, {'photo.png', 'banana.pptx'})
        self.assertTrue(all(abs(score - 1.0) < 1e-6 for info, score in results))
        self.query_encoder.encode_text.assert_called_once_with('yellow image')
        self.embedder.model.encode.assert_not_called()

    def test_broken_file_is_reported_without_stopping_other_files(self):
        Image.new('RGB', (8, 8)).save(self.folder / 'valid.png')
        (self.folder / 'empty.pdf').write_bytes(b'')
        documents, failures = self.service.analyze(self.folder)
        self.assertEqual([d.file_name for d in documents], ['valid.png'])
        self.assertEqual([p.name for p, reason in failures], ['empty.pdf'])

    def test_model_error_is_not_saved_as_success(self):
        Image.new('RGB', (8, 8)).save(self.folder / 'photo.png')
        self.image_encoder.encode.side_effect = RuntimeError('model failed')
        self.repository.save = Mock(wraps=self.repository.save)
        with self.assertRaisesRegex(RuntimeError, 'model failed'):
            self.service.analyze(self.folder)
        self.repository.save.assert_not_called()

    def test_empty_query_and_invalid_mode(self):
        self.assertEqual(self.service.search('   '), [])
        with self.assertRaises(ValueError):
            self.service.search('query', mode='unknown')

