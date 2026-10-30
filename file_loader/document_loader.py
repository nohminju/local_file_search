import os
import io
import urllib.parse
from dataclasses import dataclass
from typing import Generator, List, Optional
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation


@dataclass
class DocumentChunk:
    file_name: str
    file_path: str
    page_number: int
    text: str
    chunk_type: str
    image: Optional[Image.Image] = None


class DocumentScanner:
    SUPPORTED_EXTENSIONS = {'.pdf', '.pptx', '.jpg', '.jpeg', '.png', '.webp'}

    def __init__(self, root_dir: str):
        self.root_dir = root_dir

    def scan_files(self) -> Generator[str, None, None]:
        for root, _, files in os.walk(self.root_dir):
            for file in files:
                if file.startswith('~$'):
                    continue
                ext = os.path.splitext(file)[1].lower()
                if ext in self.SUPPORTED_EXTENSIONS:
                    yield os.path.join(root, file)


class DocumentLoader:
    def _get_clean_filename(self, file_path: str) -> str:
        raw_name = os.path.basename(file_path)
        return urllib.parse.unquote(raw_name)

    def load_image(self, file_path: str) -> List[DocumentChunk]:
        file_name = self._get_clean_filename(file_path)
        try:
            with Image.open(file_path) as img:
                pil_img = img.convert("RGB")
                return [DocumentChunk(
                    file_name=file_name,
                    file_path=file_path,
                    page_number=1,
                    text="",
                    chunk_type="image",
                    image=pil_img
                )]
        except Exception:
            return []

    def load_pdf(self, file_path: str) -> List[DocumentChunk]:
        chunks: List[DocumentChunk] = []
        file_name = self._get_clean_filename(file_path)

        try:
            reader = PdfReader(file_path)
            for page_idx, page in enumerate(reader.pages, start=1):
                try:
                    text = page.extract_text() or ""
                    clean_text = text.strip()
                    if clean_text:
                        chunks.append(DocumentChunk(
                            file_name=file_name,
                            file_path=file_path,
                            page_number=page_idx,
                            text=clean_text,
                            chunk_type="pdf"
                        ))
                except Exception:
                    continue
        except Exception:
            return []

        return chunks

    def load_pptx(self, file_path: str) -> List[DocumentChunk]:
        chunks: List[DocumentChunk] = []
        file_name = self._get_clean_filename(file_path)

        try:
            prs = Presentation(file_path)
            for slide_idx, slide in enumerate(prs.slides, start=1):
                slide_texts: List[str] = []
                extracted_image: Optional[Image.Image] = None

                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            line = paragraph.text.strip()
                            if line:
                                slide_texts.append(line)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            for cell in row.cells:
                                cell_text = cell.text.strip()
                                if cell_text:
                                    slide_texts.append(cell_text)
                    elif shape.shape_type == 13:
                        try:
                            image_bytes = shape.image.blob
                            extracted_image = Image.open(io.BytesIO(image_bytes)).convert("RGB")
                        except Exception:
                            pass

                combined_text = "\n".join(slide_texts).strip()
                if combined_text or extracted_image:
                    chunks.append(DocumentChunk(
                        file_name=file_name,
                        file_path=file_path,
                        page_number=slide_idx,
                        text=combined_text,
                        chunk_type="pptx",
                        image=extracted_image
                    ))
        except Exception:
            return []

        return chunks

    def process_file(self, file_path: str) -> List[DocumentChunk]:
        if not os.path.exists(file_path) or os.path.getsize(file_path) == 0:
            return []

        ext = os.path.splitext(file_path)[1].lower()
        if ext == '.pdf':
            return self.load_pdf(file_path)
        elif ext == '.pptx':
            return self.load_pptx(file_path)
        elif ext in {'.jpg', '.jpeg', '.png', '.webp'}:
            return self.load_image(file_path)
        return []
